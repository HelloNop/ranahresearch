import io
import re
import textwrap
import uuid
from dataclasses import dataclass

from docx import Document
from ranah_domain.enums import CitationStyle, SectionType
from ranah_domain.models.manuscript import ManuscriptSection
from ranah_domain.models.work import WorkMetadataObservation, WorkRecord
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

CITATION_TOKEN = re.compile(r"\{cite:([0-9a-fA-F-]{36})\}")


@dataclass(frozen=True, slots=True)
class Citation:
    work_id: uuid.UUID
    title: str
    authors: tuple[str, ...]
    year: int | None
    journal: str | None
    volume: str | None
    issue: str | None
    pages: str | None
    doi: str | None
    incomplete: tuple[str, ...]


async def citation_for_work(session: AsyncSession, work_id: uuid.UUID) -> Citation:
    work = await session.get(WorkRecord, work_id)
    if work is None:
        raise ValueError(f"Unknown canonical work {work_id}")
    observation = await session.scalar(
        select(WorkMetadataObservation)
        .where(
            WorkMetadataObservation.work_id == work_id,
            WorkMetadataObservation.field_name == "authors",
        )
        .order_by(WorkMetadataObservation.observed_at.desc())
        .limit(1)
    )
    raw = observation.field_value if observation else {}
    values = raw.get("value", raw.get("authors", [])) if isinstance(raw, dict) else []
    if not isinstance(values, list):
        values = []
    authors = tuple(
        str(item.get("name") or item.get("display_name") or "").strip()
        if isinstance(item, dict)
        else str(item).strip()
        for item in values
        if item
    )
    authors = tuple(author for author in authors if author)
    incomplete = tuple(
        name
        for name, value in (
            ("authors", authors),
            ("year", work.publication_year),
            ("title", work.title),
            ("journal", work.journal),
            ("doi", work.doi),
        )
        if not value
    )
    return Citation(
        work_id=work.id,
        title=work.title,
        authors=authors,
        year=work.publication_year,
        journal=work.journal,
        volume=work.volume,
        issue=work.issue,
        pages=work.pages,
        doi=work.doi,
        incomplete=incomplete,
    )


def in_text(citation: Citation, style: CitationStyle, number: int) -> str:
    if style == CitationStyle.VANCOUVER:
        return f"[{number}]"
    surnames = [author.split()[-1] for author in citation.authors]
    if not surnames:
        author_text = citation.title[:40]
    elif len(surnames) == 1:
        author_text = surnames[0]
    elif len(surnames) == 2:
        author_text = f"{surnames[0]} & {surnames[1]}"
    else:
        author_text = f"{surnames[0]} et al."
    return f"({author_text}, {citation.year or 'n.d.'})"


def bibliography_entry(citation: Citation, style: CitationStyle, number: int) -> str:
    authors = ", ".join(citation.authors) or "Author unavailable"
    year = str(citation.year) if citation.year else "n.d."
    publication = ". ".join(
        item
        for item in (
            citation.journal,
            citation.volume + (f"({citation.issue})" if citation.issue else "")
            if citation.volume
            else None,
            citation.pages,
        )
        if item
    )
    doi = f" https://doi.org/{citation.doi}" if citation.doi else ""
    if style == CitationStyle.VANCOUVER:
        return f"{number}. {authors}. {citation.title}. {publication}. {year}.{doi}".replace(
            ". .", "."
        )
    return f"{authors} ({year}). {citation.title}. {publication}.{doi}".replace(". .", ".")


async def resolve_sections(
    session: AsyncSession, sections: list[ManuscriptSection], style: CitationStyle
) -> tuple[list[tuple[ManuscriptSection, str]], list[str], list[dict[str, object]]]:
    ordered_ids: list[uuid.UUID] = []
    for section in sections:
        for match in CITATION_TOKEN.finditer(section.content):
            work_id = uuid.UUID(match.group(1))
            if work_id not in ordered_ids:
                ordered_ids.append(work_id)
    citations = [await citation_for_work(session, work_id) for work_id in ordered_ids]
    if style == CitationStyle.APA_7:
        citations.sort(
            key=lambda row: ((row.authors[0] if row.authors else row.title).lower(), row.year or 0)
        )
    numbers = {row.work_id: index for index, row in enumerate(citations, 1)}
    rendered: list[tuple[ManuscriptSection, str]] = []
    for section in sections:
        content = CITATION_TOKEN.sub(
            lambda match: in_text(
                next(row for row in citations if row.work_id == uuid.UUID(match.group(1))),
                style,
                numbers[uuid.UUID(match.group(1))],
            ),
            section.content,
        )
        rendered.append((section, content))
    bibliography = [bibliography_entry(row, style, numbers[row.work_id]) for row in citations]
    warnings: list[dict[str, object]] = [
        {"work_id": str(row.work_id), "missing": list(row.incomplete)}
        for row in citations
        if row.incomplete
    ]
    return rendered, bibliography, warnings


def markdown(
    title: str, sections: list[tuple[ManuscriptSection, str]], references: list[str]
) -> bytes:
    parts = [f"# {title}"]
    for section, content in sections:
        if section.section_type not in {SectionType.TITLE, SectionType.REFERENCES}:
            parts.extend((f"## {section.heading}", content))
    if references:
        parts.extend(("## References", "\n".join(f"- {row}" for row in references)))
    return "\n\n".join(parts).encode()


def latex(
    title: str, sections: list[tuple[ManuscriptSection, str]], references: list[str]
) -> bytes:
    def escape(value: str) -> str:
        return re.sub(r"([#$%&_{}])", r"\\\1", value).replace("~", r"\textasciitilde{}")

    body = [
        r"\documentclass{article}",
        r"\usepackage[utf8]{inputenc}",
        f"\\title{{{escape(title)}}}",
        r"\begin{document}",
        r"\maketitle",
    ]
    for section, content in sections:
        if section.section_type not in {SectionType.TITLE, SectionType.REFERENCES}:
            body.extend((f"\\section{{{escape(section.heading)}}}", escape(content)))
    if references:
        body.extend((r"\section{References}", r"\begin{enumerate}"))
        body.extend(f"\\item {escape(row)}" for row in references)
        body.append(r"\end{enumerate}")
    body.append(r"\end{document}")
    return "\n".join(body).encode()


def docx(title: str, sections: list[tuple[ManuscriptSection, str]], references: list[str]) -> bytes:
    document = Document()
    document.add_heading(title, 0)
    for section, content in sections:
        if section.section_type not in {SectionType.TITLE, SectionType.REFERENCES}:
            document.add_heading(section.heading, level=1)
            document.add_paragraph(content)
    if references:
        document.add_heading("References", level=1)
        for row in references:
            document.add_paragraph(row)
    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def pdf(title: str, sections: list[tuple[ManuscriptSection, str]], references: list[str]) -> bytes:
    lines = [title, ""]
    for section, content in sections:
        if section.section_type not in {SectionType.TITLE, SectionType.REFERENCES}:
            lines.extend((section.heading, *textwrap.wrap(content, 92), ""))
    if references:
        lines.extend(
            ("References", *[f"{index}. {row}" for index, row in enumerate(references, 1)])
        )

    pages = [lines[index : index + 48] for index in range(0, len(lines), 48)] or [[title]]
    objects: list[bytes] = [
        b"",
        b"",
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Times-Roman >>",
    ]
    page_ids: list[int] = []
    for page_lines in pages:
        page_id = len(objects) + 1
        content_id = page_id + 1
        page_ids.append(page_id)
        objects.append(
            f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            f"/Resources << /Font << /F1 3 0 R >> >> /Contents {content_id} 0 R >>".encode()
        )
        commands = ["BT /F1 11 Tf 14 TL 60 740 Td"]
        for line in page_lines:
            safe = line.encode("latin-1", "replace").decode("latin-1")
            safe = safe.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")
            commands.append(f"({safe}) Tj T*")
        commands.append("ET")
        stream = "\n".join(commands).encode("latin-1")
        objects.append(f"<< /Length {len(stream)} >>\nstream\n".encode() + stream + b"\nendstream")
    objects[0] = b"<< /Type /Catalog /Pages 2 0 R >>"
    kids = " ".join(f"{page_id} 0 R" for page_id in page_ids)
    objects[1] = f"<< /Type /Pages /Kids [{kids}] /Count {len(page_ids)} >>".encode()

    output = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for number, body in enumerate(objects, 1):
        offsets.append(len(output))
        output.extend(f"{number} 0 obj\n".encode() + body + b"\nendobj\n")
    xref = len(output)
    output.extend(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    output.extend(b"".join(f"{offset:010d} 00000 n \n".encode() for offset in offsets[1:]))
    output.extend(
        f"trailer << /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF".encode()
    )
    return bytes(output)
